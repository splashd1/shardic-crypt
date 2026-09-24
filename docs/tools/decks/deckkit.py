"""Shared visual system for the shardic v4.1 deck set (matches shardic_deck_v3_2.pptx)."""
import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

_HERE = os.path.dirname(os.path.abspath(__file__))
ICONS = os.path.join(_HERE, "icons") + os.sep          # white icon PNGs reused from shardic_deck_v3_2.pptx
MEDIA = os.path.join(_HERE, "..", "..", "media") + os.sep  # docs/media diagrams
IC = dict(warn="image-12-2.png", userclock="image-13-2.png", tree="image-13-3.png", key="image-2-1.png",
          cubes="image-3-1.png", check="image-3-7.png", shield="image-4-1.png", users="image-4-2.png",
          bank="image-4-3.png", lock="image-7-2.png", usershield="image-8-1.png", scales="image-9-9.png",
          case="image10.png", puzzle="image12.png", db="image14.png", mail="image31.png", bell="image34.png")

def C(h): return RGBColor.from_string(h)
NAVY, DEEP, BLUE = "24316E", "141B4D", "263874"
GOLD, GOLD_D, ICE, SOFT = "E7B93E", "9C6F0B", "CADCFC", "AEB9DC"
MUTED, LIGHT, GREEN, BROWN, WHITE = "5B6B8C", "EEF2FB", "3EBD6E", "3A2F14", "FFFFFF"
BODY_FONT, CODE_FONT = "Calibri", "Consolas"
W, H = 13.333, 7.5


class Deck:
    def __init__(self, label):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.blank = self.prs.slide_layouts[6]
        self.label = label
        self.n = 0

    # ---- slide scaffolding -------------------------------------------------
    def slide(self, dark=False, eyebrow=None, title=None, subtitle=None, notes=None):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        fill = s.background.fill
        fill.solid()
        fill.fore_color.rgb = C(DEEP if dark else WHITE)
        if eyebrow:
            text(s, 0.6, 0.38, 9, 0.3, eyebrow.upper(), 11, GOLD if dark else GOLD_D, bold=True, spacing=100)
        if title:
            text(s, 0.6, 0.62, 12.1, 0.75, title, 30, WHITE if dark else DEEP, bold=True)
        if subtitle:
            text(s, 0.6, 1.3, 12.1, 0.6, subtitle, 15, SOFT if dark else MUTED)
        if not dark:
            text(s, 0.6, 7.02, 6, 0.28, "SHARDIC  ·  " + self.label, 9, MUTED, bold=True, spacing=60)
            text(s, 11.9, 7.02, 0.83, 0.28, str(self.n), 9, MUTED, align="r")
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def save(self, path):
        self.prs.save(path)


# ---- primitives ---------------------------------------------------------------
def _runs(p, s, size, color, bold, font, italic):
    # **bold** inline markup
    for i, part in enumerate(re.split(r"\*\*", s)):
        if not part:
            continue
        r = p.add_run()
        r.text = part
        f = r.font
        f.size, f.name = Pt(size), font
        f.bold = bold or (i % 2 == 1)
        f.italic = italic
        f.color.rgb = C(color)


def text(s, x, y, w, h, content, size=14, color=MUTED, bold=False, font=BODY_FONT, align="l",
         anchor="t", italic=False, bullets=False, gap=6, spacing=None, line=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = dict(t=MSO_ANCHOR.TOP, m=MSO_ANCHOR.MIDDLE, b=MSO_ANCHOR.BOTTOM)[anchor]
    items = content if isinstance(content, list) else [content]
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = dict(l=PP_ALIGN.LEFT, c=PP_ALIGN.CENTER, r=PP_ALIGN.RIGHT)[align]
        if i:
            p.space_before = Pt(gap)
        if line:
            p.line_spacing = line
        _runs(p, item, size, color, bold, font, italic)
        if spacing:
            for r in p.runs:
                r.font._element.set("spc", str(spacing))
        if bullets:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Emu(Inches(0.22)))); pPr.set("indent", str(-Emu(Inches(0.22))))
            bc = pPr.makeelement(qn("a:buClr"), {}); sc = bc.makeelement(qn("a:srgbClr"), {"val": GOLD})
            bc.append(sc); pPr.append(bc)
            bu = pPr.makeelement(qn("a:buChar"), {"char": "•"}); pPr.append(bu)
    return tb


def rect(s, x, y, w, h, fill=NAVY, radius=0.06, line=None, shape=None):
    shp = s.shapes.add_shape(shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE),
                             Inches(x), Inches(y), Inches(w), Inches(h))
    if radius and shape is None:
        shp.adjustments[0] = radius
    shp.fill.solid(); shp.fill.fore_color.rgb = C(fill)
    if line:
        shp.line.color.rgb = C(line); shp.line.width = Pt(1.25)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def icon(s, x, y, name, d=0.62, fill=BLUE):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = C(fill); c.line.fill.background(); c.shadow.inherit = False
    pad = d * 0.24
    s.shapes.add_picture(ICONS + IC[name], Inches(x + pad), Inches(y + pad), Inches(d - 2 * pad), Inches(d - 2 * pad))


def badge(s, x, y, label, d=0.62, fill=GOLD, color=DEEP, size=20):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = C(fill); c.line.fill.background(); c.shadow.inherit = False
    text(s, x, y, d, d, label, size, color, bold=True, align="c", anchor="m")


def pill(s, x, y, label, fill=GREEN, color=DEEP, w=None, size=10):
    w = w or (0.2 + 0.085 * len(label))
    rect(s, x, y, w, 0.3, fill, radius=0.5)
    text(s, x, y, w, 0.3, label.upper(), size, color, bold=True, align="c", anchor="m", spacing=60)
    return w


def card(s, x, y, w, h, head=None, body=None, ic=None, fill=NAVY, head_size=16, body_size=12.5, tag=None, tagfill=GREEN, hl=1):
    rect(s, x, y, w, h, fill)
    ty = y + 0.28
    if ic:
        icon(s, x + 0.3, y + 0.3, ic, 0.56)
        ty = y + 1.02
    if tag:
        pill(s, x + w - 0.3 - (0.2 + 0.085 * len(tag)), y + 0.3, tag, tagfill)
    if head:
        text(s, x + 0.3, ty, w - 0.6, 0.5, head, head_size, WHITE, bold=True)
        ty += 0.14 + head_size / 72 * 1.25 * hl
    if body:
        text(s, x + 0.3, ty, w - 0.6, y + h - ty - 0.2, body, body_size, ICE, bullets=isinstance(body, list), gap=5)


def code(s, x, y, w, h, lines, size=12, fill=DEEP):
    rect(s, x, y, w, h, fill, radius=0.03)
    text(s, x + 0.3, y + 0.25, w - 0.6, h - 0.5, lines, size, "D6DEF5", font=CODE_FONT, gap=1)


def callout(s, x, y, w, h, content, size=15, fill=BLUE, color=WHITE, italic=False):
    rect(s, x, y, w, h, fill, radius=0.08)
    text(s, x + 0.35, y, w - 0.7, h, content, size, color, anchor="m", italic=italic)


def picture(s, name, x, y, w=None, h=None, frame=True):
    from PIL import Image
    iw, ih = Image.open(MEDIA + name).size
    if w and not h: h = w * ih / iw
    if h and not w: w = h * iw / ih
    if frame:
        rect(s, x - 0.15, y - 0.15, w + 0.3, h + 0.3, LIGHT, radius=0.04)
    s.shapes.add_picture(MEDIA + name, Inches(x), Inches(y), Inches(w), Inches(h))
    return w, h


def stat(s, x, y, w, big, small, color=DEEP, sub=MUTED, size=60):
    text(s, x, y, w, 1.1, big, size, color, bold=True)
    text(s, x, y + size / 72 * 1.25, w, 0.9, small, 14, sub)
